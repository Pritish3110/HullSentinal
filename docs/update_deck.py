#!/usr/bin/env python3
"""
Extends the HullSentinal project deck with the Introduction / Literature Survey /
Problem Statement / Proposed Work / Results and Discussion / Conclusion / References
structure, reusing the original theme tokens exactly.

Usage:  python3 docs/update_deck.py
"""

import os
import copy

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), '..'))
DECK = os.path.join(ROOT, 'docs', 'Hull_Cleaning_Cognitive_Project.pptx')
FIGS = os.path.join(ROOT, 'docs', 'paper_figures')

# ---- theme tokens lifted from the original deck ----------------------------
NAVY = RGBColor(0x21, 0x29, 0x5C)
DEEP = RGBColor(0x06, 0x5A, 0x82)
MID = RGBColor(0x1C, 0x72, 0x93)
TEAL = RGBColor(0x0A, 0x93, 0x96)
INK = RGBColor(0x1B, 0x1B, 0x1B)
MUTED = RGBColor(0x5B, 0x6B, 0x75)
PANEL = RGBColor(0xF4, 0xF8, 0xFA)
PALE = RGBColor(0xBF, 0xD9, 0xE6)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
CRIMSON = RGBColor(0xC9, 0x28, 0x3E)

TITLE_FONT = 'Cambria'
BODY_FONT = 'Calibri'


def textbox(slide, x, y, w, h, anchor=MSO_ANCHOR.TOP):
    sh = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = sh.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    return sh, tf


def write(tf, blocks, first=True):
    """blocks: list of dicts {t, size, bold, color, font, align, space_before}."""
    for b in blocks:
        para = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        para.alignment = b.get('align', PP_ALIGN.LEFT)
        if b.get('space_before'):
            para.space_before = Pt(b['space_before'])
        if b.get('space_after'):
            para.space_after = Pt(b['space_after'])
        for seg in (b['t'] if isinstance(b['t'], list) else [{'t': b['t']}]):
            run = para.add_run()
            run.text = seg['t']
            f = run.font
            f.name = seg.get('font', b.get('font', BODY_FONT))
            f.size = Pt(seg.get('size', b.get('size', 12)))
            f.bold = seg.get('bold', b.get('bold', False))
            f.italic = seg.get('italic', b.get('italic', False))
            f.color.rgb = seg.get('color', b.get('color', INK))
            if seg.get('sub'):
                f._rPr.set('baseline', '-25000')
            elif seg.get('sup'):
                f._rPr.set('baseline', '30000')
    return tf


def title(slide, text):
    sh, tf = textbox(slide, 0.60, 0.40, 12.13, 0.90, MSO_ANCHOR.MIDDLE)
    write(tf, [{'t': text, 'size': 32, 'bold': True, 'color': NAVY, 'font': TITLE_FONT}])
    return sh


def subtitle(slide, text):
    sh, tf = textbox(slide, 0.60, 1.15, 11.50, 0.40)
    write(tf, [{'t': text, 'size': 12.5, 'color': MUTED}])
    return sh


def card(slide, x, y, w, h, fill=PANEL, radius=0.055):
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    sh.line.fill.background()
    sh.shadow.inherit = False
    try:
        sh.adjustments[0] = radius
    except Exception:
        pass
    sh.text_frame.text = ''
    return sh


def bar(slide, x, y, w, h, fill):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,
                                Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    sh.line.fill.background()
    sh.shadow.inherit = False
    sh.text_frame.text = ''
    return sh


def pagenum(slide, n):
    sh, tf = textbox(slide, 12.60, 7.05, 0.50, 0.30)
    write(tf, [{'t': str(n), 'size': 10, 'color': MUTED, 'align': PP_ALIGN.RIGHT}])


def blank(prs):
    return prs.slides.add_slide(prs.slide_layouts[0])


def grid(slide, x0, y0, col_w, header, rows, body_sizes=(11.0,), row_h=0.42,
         head_h=0.46, aligns=None, highlight_last=False):
    """Theme-consistent table: deep header bar, alternating white/panel rows."""
    aligns = aligns or [PP_ALIGN.LEFT] * len(col_w)
    bar(slide, x0, y0, sum(col_w), head_h, DEEP)
    cx = x0
    for i, (c, w) in enumerate(zip(header, col_w)):
        _, tf = textbox(slide, cx + 0.12, y0 + 0.06, w - 0.24, head_h - 0.12, MSO_ANCHOR.MIDDLE)
        write(tf, [{'t': c, 'size': 11.5, 'bold': True, 'color': WHITE, 'align': aligns[i]}])
        cx += w
    y = y0 + head_h
    for ri, row in enumerate(rows):
        bar(slide, x0, y, sum(col_w), row_h, WHITE if ri % 2 == 0 else PANEL)
        cx = x0
        for i, (c, w) in enumerate(zip(row, col_w)):
            _, tf = textbox(slide, cx + 0.12, y + 0.04, w - 0.24, row_h - 0.08, MSO_ANCHOR.MIDDLE)
            bold = highlight_last and (ri == len(rows) - 1)
            col = DEEP if bold else INK
            write(tf, [{'t': c, 'size': body_sizes[0], 'bold': bold, 'color': col,
                        'align': aligns[i]}])
            cx += w
        y += row_h
    return y


def picture(slide, fname, x, y, w):
    path = os.path.join(FIGS, fname)
    return slide.shapes.add_picture(path, Inches(x), Inches(y), width=Inches(w))


def picture_fit(slide, fname, cx, y, max_w, max_h):
    """Places the image centred on cx, scaled to fit inside max_w x max_h."""
    from PIL import Image as _Img
    path = os.path.join(FIGS, fname)
    with _Img.open(path) as im:
        iw, ih = im.size
    w = max_w
    h = w * ih / iw
    if h > max_h:
        h = max_h
        w = h * iw / ih
    return slide.shapes.add_picture(path, Inches(cx - w / 2), Inches(y),
                                    width=Inches(w), height=Inches(h)), h


def caption(slide, x, y, w, text):
    _, tf = textbox(slide, x, y, w, 0.34)
    write(tf, [{'t': text, 'size': 9.5, 'color': MUTED, 'align': PP_ALIGN.CENTER}])


# ===========================================================================
#  New slides
# ===========================================================================
def slide_problem_statement(prs):
    s = blank(prs)
    title(s, 'Problem Statement')
    subtitle(s, 'What to clean first, in what order, and how to prove it was actually cleaned')

    card(s, 0.60, 1.70, 12.10, 1.05)
    _, tf = textbox(s, 0.95, 1.88, 11.40, 0.75, MSO_ANCHOR.MIDDLE)
    write(tf, [{'t': 'Given hull defects of differing type, severity and detection confidence, '
                     'how should a two-robot system order remediation and confirm completion — '
                     'without a fixed schedule or a blind task-done assumption?',
                'size': 13, 'italic': True, 'color': NAVY}])

    items = [
        ('Risk blindness',
         'Dispatch minimises travel distance only. Defect class τ(d) and severity S(d) never enter '
         'the objective, so soft biofouling at 2 m outranks a structural crack at 6 m.'),
        ('Open-loop completion',
         'Arrival at the commanded pose is treated as success. A partial cavitation pass is '
         'indistinguishable from a complete one, and the residual defect is never revisited.'),
        ('Static criticality',
         'Severity is sampled once at detection. A crack that propagates during the mission is '
         'scheduled on stale information and can cross the critical threshold unnoticed.'),
    ]
    x = 0.60
    for name, desc in items:
        card(s, x, 3.05, 3.95, 2.25)
        bar(s, x, 3.05, 3.95, 0.07, MID)
        _, tf = textbox(s, x + 0.30, 3.35, 3.35, 0.40)
        write(tf, [{'t': name, 'size': 14, 'bold': True, 'color': NAVY}])
        _, tf = textbox(s, x + 0.30, 3.85, 3.35, 1.55)
        write(tf, [{'t': desc, 'size': 11.5, 'color': MUTED}])
        x += 4.15

    card(s, 0.60, 5.90, 12.10, 0.95, PALE)
    _, tf = textbox(s, 0.95, 6.05, 11.40, 0.65, MSO_ANCHOR.MIDDLE)
    write(tf, [{'t': 'Objective:  maximise rank correlation rₛ between the realised servicing order '
                     'and ground-truth structural risk, subject to zero geofence excursions and zero '
                     'inter-robot collisions — and retire no defect without physical evidence.',
                'size': 12.5, 'bold': True, 'color': NAVY}])
    return s


def slide_utility(prs):
    s = blank(prs)
    title(s, 'Proposed Work: Utility Function & Core Algorithms')
    subtitle(s, 'Explicit, interpretable prioritisation over true geodesic cost')

    card(s, 0.60, 1.70, 12.10, 1.15, PANEL)
    _, tf = textbox(s, 0.95, 1.88, 11.40, 0.80, MSO_ANCHOR.MIDDLE)
    F = lambda t, **k: dict(t=t, size=17, bold=True, color=DEEP, font=TITLE_FONT, **k)
    write(tf, [{'t': [
        F('U(d)  =  w'), F('1', sub=True), F(' · Ŝ(d) · W'), F('risk', sub=True),
        F('(τ(d)) · Ĉ(d)    −   α · g'), F('A*', sub=True),
        F('(p'), F('c', sub=True), F(' , p'), F('d', sub=True),
        F(')    +   β · 𝟙'), F('requeued', sub=True),
        F('    +   Ω · 𝟙'), F('critical', sub=True),
    ], 'align': PP_ALIGN.CENTER}])

    terms = [
        ([{'t': 'Ŝ · W'}, {'t': 'risk', 'sub': True}, {'t': ' · Ĉ'}], 'Expected structural consequence: severity scaled by class '
         'risk, discounted by detection confidence.'),
        ([{'t': '− α · g'}, {'t': 'A*', 'sub': True}], 'True 8-connected A* geodesic distance with 0.55 m obstacle '
         'inflation — not Euclidean distance.'),
        ('+ β', 'Escalation bonus applied when verification re-queues a defect.'),
        ('+ Ω', 'Emergency override, set an order of magnitude above the first term so it acts '
         'lexicographically.'),
    ]
    x = 0.60
    for sym, desc in terms:
        card(s, x, 3.15, 2.95, 1.85)
        _, tf = textbox(s, x + 0.22, 3.35, 2.50, 0.42)
        write(tf, [{'t': sym, 'size': 14, 'bold': True, 'color': TEAL, 'font': TITLE_FONT}])
        _, tf = textbox(s, x + 0.22, 3.82, 2.50, 1.05)
        write(tf, [{'t': desc, 'size': 10.5, 'color': MUTED}])
        x += 3.10

    rows = [
        [[{'t': 'Risk weight W'}, {'t': 'risk', 'sub': True}], 'crack 1.00  >  corrosion 0.70  >  biofouling 0.40'],
        ['Belief fusion', 'Ŝₖ = (1−λ)Ŝₖ₋₁ + λ zₖ ,  λ = 0.50;   Ĉₖ = max(Ĉₖ₋₁ , cₖ)'],
        ['Degradation', 'dS/dt = γτ  ⇒  deferred cracks genuinely propagate (γ = 0.0080 s⁻¹)'],
        ['Lane gate', [{'t': 'E(d) = [ ℓ(d) < ℓ'}, {'t': 'map', 'sub': True}, {'t': ' ]  ∨  [ ρ(Ŝ(d)) = CRITICAL ]'}]],
        ['Verification', 'residual ≤ θᵥ (0.08) ⇒ VERIFIED_CLEAN,  else RE-QUEUED with bonus β'],
    ]
    grid(s, 0.60, 5.25, [3.10, 9.00], ['Mechanism', 'Definition'], rows, row_h=0.36, head_h=0.40)
    return s


def slide_results_walkthrough(prs):
    s = blank(prs)
    title(s, 'Results: Mission Walkthrough')
    subtitle(s, 'One complete 14-defect mission — lane gating, risk ordering, preemption, verification')
    _, fh = picture_fit(s, 'fig4_stages.png', 6.665, 1.62, 10.40, 4.42)
    caption(s, 0.95, 1.62 + fh + 0.10, 11.40,
            'Simulator output at four mission stages. Grey = UNDETECTED, orange = QUEUED, '
            'red = CLEANING, green = VERIFIED_CLEAN. Cyan trace = mapper serpentine; orange trace = cleaner tour.')
    card(s, 0.60, 6.74, 12.10, 0.62, PALE)
    _, tf = textbox(s, 0.95, 6.84, 11.40, 0.42, MSO_ANCHOR.MIDDLE)
    write(tf, [{'t': 'Cleaner holds until Lane 1 is released → services crack, then corrosion, then '
                     'biofouling → preempts for an escalating crack via the context stack → resumes '
                     '→ all 14 defects verified clean.',
                'size': 11.5, 'bold': True, 'color': NAVY}])
    return s


def slide_results_table(prs):
    s = blank(prs)
    title(s, 'Results: Quantitative Comparison')
    subtitle(s, '60 independent missions — 6 scenarios × 5 seeds × 2 architectures')

    rows = [
        ['nominal', '85.3 ± 2.1', '49.7 ± 1.5', '+0.77', '−0.07', '0 / 0'],
        ['swarm_scaling', '60.6 ± 1.5', '53.9 ± 1.4', '+0.23', '+0.05', '0 / 0'],
        ['obstacles', '89.0 ± 1.5', '49.7 ± 1.5', '+0.83', '−0.07', '0 / 0'],
        ['high_sensor_noise', '85.3 ± 2.1', '49.7 ± 1.5', '+0.77', '−0.07', '0 / 0'],
        ['agent_failure', '87.6 ± 2.5', '49.7 ± 1.5', '+0.81', '−0.07', '0 / 0'],
        ['marine_current', '85.3 ± 2.1', '49.7 ± 1.5', '+0.77', '−0.07', '0 / 0'],
        ['Overall (n = 30 each)', '82.17', '50.38', '+0.70', '−0.05', '0 / 0'],
    ]
    al = [PP_ALIGN.LEFT] + [PP_ALIGN.CENTER] * 5
    grid(s, 0.60, 1.72, [2.90, 2.05, 2.05, 1.75, 1.75, 1.60],
         ['Scenario', 'Order corr. % (Ours)', 'Order corr. % (FIFO)',
          'rₛ (Ours)', 'rₛ (FIFO)', 'OOB / Coll.'],
         rows, row_h=0.40, head_h=0.52, aligns=al, highlight_last=True)

    stats = [
        ('+31.8 pp', 'ordering correctness gain over the FIFO baseline'),
        ('d = 4.37', 'Cohen’s d; Welch t = 16.91 — a decisive effect'),
        ('36%', 'of remediations caught and re-queued by verification'),
        ('0', 'geofence, lane and collision violations in 60 missions'),
    ]
    x = 0.60
    for big, desc in stats:
        card(s, x, 5.15, 2.95, 1.55)
        _, tf = textbox(s, x + 0.22, 5.35, 2.50, 0.52)
        write(tf, [{'t': big, 'size': 24, 'bold': True, 'color': TEAL, 'font': TITLE_FONT}])
        _, tf = textbox(s, x + 0.22, 5.92, 2.50, 0.70)
        write(tf, [{'t': desc, 'size': 10.5, 'color': MUTED}])
        x += 3.10

    _, tf = textbox(s, 0.60, 6.85, 12.10, 0.32)
    write(tf, [{'t': 'Both architectures retire all 14 defects in every mission; the ordering gain is '
                     'not bought by sacrificing completeness.',
                'size': 10.5, 'italic': True, 'color': MUTED}])
    return s


def slide_results_charts(prs):
    s = blank(prs)
    title(s, 'Results: Ordering Quality & Risk Alignment')
    subtitle(s, 'The baseline is pinned near chance in every scenario; the proposed system is not')
    _, h1 = picture_fit(s, 'fig5_order.png', 3.575, 1.72, 5.55, 3.30)
    _, h2 = picture_fit(s, 'fig6_correlation.png', 9.725, 1.72, 5.55, 3.30)
    cy = 1.72 + max(h1, h2) + 0.12
    caption(s, 0.70, cy, 5.75, 'Defect-ordering correctness by scenario (mean ± SD, 5 seeds).')
    caption(s, 6.85, cy, 5.75, 'Spearman rₛ between servicing order and structural risk.')
    card(s, 0.60, 6.10, 12.10, 1.00, PANEL)
    _, tf = textbox(s, 0.95, 6.24, 11.40, 0.72, MSO_ANCHOR.MIDDLE)
    write(tf, [{'t': 'The single weak case is swarm_scaling (60.6%): three cleaners draw concurrently '
                     'from one ranked queue, so completion order is partly governed by travel time. '
                     'The remedy is architectural — replace the shared queue with a consensus-based '
                     'bundle auction (CBBA) in which each cleaner bids on its own marginal geodesic cost.',
                'size': 11.5, 'color': INK}])
    return s


def slide_discussion(prs):
    s = blank(prs)
    title(s, 'Discussion: Trade-offs & Limitations')
    subtitle(s, 'What the prioritisation costs, and where the claims stop')

    card(s, 0.60, 1.70, 5.95, 2.45)
    _, tf = textbox(s, 0.90, 1.92, 5.35, 0.40)
    write(tf, [{'t': 'The cost of prioritisation', 'size': 14.5, 'bold': True, 'color': NAVY}])
    _, tf = textbox(s, 0.90, 2.42, 5.35, 1.55)
    write(tf, [{'t': 'Missions take 106.7 s against 76.4 s for FIFO — a 39.7% increase. Lane gating '
                     'idles the cleaner for separation, risk ordering abandons distance optimality, and '
                     'verification adds a re-scan per defect. Each buys a property the baseline lacks, '
                     'and the coefficient α moves the operating point directly.',
                'size': 11.5, 'color': MUTED}])

    card(s, 6.75, 1.70, 5.95, 2.45, PALE)
    _, tf = textbox(s, 7.05, 1.92, 5.35, 0.40)
    write(tf, [{'t': 'Why the trade is favourable', 'size': 14.5, 'bold': True, 'color': NAVY}])
    _, tf = textbox(s, 7.05, 2.42, 5.35, 1.55)
    write(tf, [{'t': 'In the agent_failure scenario the open-loop baseline reports 100% success while '
                     'leaving residual defects on the hull. Closed-loop verification rejected 36% of '
                     'remediations there. For a task where a deferred crack is the dominant risk, '
                     '30 s is a cheap price for correct ordering and verified completion.',
                'size': 11.5, 'color': MUTED}])

    lims = [
        ('Simulation only',
         'Perception is a noisy oracle, not a learned detector; noise robustness characterises the '
         'belief filter, not a vision pipeline.'),
        ('Planar manifold',
         'Real hulls are doubly curved. Geodesic planning over a triangulated mesh is required before '
         'path costs transfer.'),
        ('Single-server metric',
         'Order correctness assumes one sequential cleaner and therefore under-reports multi-cleaner '
         'performance.'),
    ]
    x = 0.60
    for name, desc in lims:
        card(s, x, 4.40, 3.95, 2.00)
        bar(s, x, 4.40, 3.95, 0.07, MID)
        _, tf = textbox(s, x + 0.28, 4.66, 3.40, 0.38)
        write(tf, [{'t': name, 'size': 13, 'bold': True, 'color': NAVY}])
        _, tf = textbox(s, x + 0.28, 5.10, 3.40, 1.15)
        write(tf, [{'t': desc, 'size': 11, 'color': MUTED}])
        x += 4.15

    _, tf = textbox(s, 0.60, 6.60, 12.10, 0.40)
    write(tf, [{'t': 'Utility weights were set by design reasoning rather than fitted to data, and are '
                     'presented as an interpretable default rather than an optimum.',
                'size': 11, 'italic': True, 'color': MUTED}])
    return s


def slide_conclusion(prs):
    s = blank(prs)
    title(s, 'Conclusion & Future Work')

    card(s, 0.60, 1.55, 12.10, 1.35, PALE)
    _, tf = textbox(s, 0.95, 1.72, 11.40, 1.00, MSO_ANCHOR.MIDDLE)
    write(tf, [{'t': 'HullSentinal replaces geometric dispatch and self-reported completion with an '
                     'explicit cognition layer: a probabilistic defect belief carrying severity and '
                     'confidence, an interpretable structural-utility ranking over true geodesic cost, '
                     'and a mandatory physical re-scan before any defect may be retired.',
                'size': 13, 'bold': True, 'color': NAVY}])

    wins = [
        ('82.17%', 'ordering correctness, from 50.38% (FIFO)'),
        ('rₛ = +0.70', 'risk alignment, from −0.05 (FIFO)'),
        ('14 / 14', 'defects verified clean every mission'),
        ('0', 'safety violations across 60 missions'),
    ]
    x = 0.60
    for big, desc in wins:
        card(s, x, 3.10, 2.95, 1.35)
        _, tf = textbox(s, x + 0.22, 3.26, 2.50, 0.50)
        write(tf, [{'t': big, 'size': 21, 'bold': True, 'color': TEAL, 'font': TITLE_FONT}])
        _, tf = textbox(s, x + 0.22, 3.80, 2.50, 0.52)
        write(tf, [{'t': desc, 'size': 10.5, 'color': MUTED}])
        x += 3.10

    _, tf = textbox(s, 0.60, 4.75, 12.10, 0.36)
    write(tf, [{'t': 'Future work', 'size': 15, 'bold': True, 'color': NAVY}])

    future = [
        ('Swarm scaling', 'Replace the shared ranked queue with a consensus-based bundle auction so '
                          'N cleaners bid on bundles by marginal geodesic cost.'),
        ('3D curved hulls', 'Replace the planar occupancy grid with geodesic planning over a '
                            'triangulated hull mesh and surface-normal adhesion.'),
        ('Real perception', 'Substitute the oracle detector with a learned segmentation model; the '
                            'cognition layer is detector-agnostic and needs no redesign.'),
    ]
    x = 0.60
    for name, desc in future:
        card(s, x, 5.20, 3.95, 1.70)
        bar(s, x, 5.20, 3.95, 0.07, TEAL)
        _, tf = textbox(s, x + 0.28, 5.44, 3.40, 0.36)
        write(tf, [{'t': name, 'size': 13, 'bold': True, 'color': DEEP}])
        _, tf = textbox(s, x + 0.28, 5.86, 3.40, 0.95)
        write(tf, [{'t': desc, 'size': 11, 'color': MUTED}])
        x += 4.15
    return s


REFERENCES = [
    'H. Su, S. Liu, L. Zhang, Y. Chen, and C. Yang, \u201cBiofouling recognition and boundary tracking '
    'control for underwater cleaning robots,\u201d Ocean Engineering, vol. 295, art. 116707, 2024.',
    '\u201cAutomated detection and classification of cracks in ship hull structures using deep '
    'learning,\u201d Ocean Engineering, 2026.',
    '\u201cDetection of cracks and corrosion for automated vessel visual inspection,\u201d LIACI dataset, '
    'EU FP7 MINOAS project.',
    '\u201cDesign, perception, planning and control for underwater cleaning robotics: progress, '
    'challenges and trends,\u201d Ocean Engineering review, 2025.',
    'M. Chwe\u0142atiuk et al., \u201cAlgorithm for determining the position of a ship hull-cleaning '
    'robot,\u201d Applied Sciences, vol. 15, no. 10, art. 5705, 2025.',
    'B. P. Gerkey and M. J. Matari\u0107, \u201cA formal analysis and taxonomy of task allocation in '
    'multi-robot systems,\u201d Int. J. Robotics Research, vol. 23, no. 9, pp. 939\u2013954, 2004.',
    'H.-L. Choi, L. Brunet, and J. P. How, \u201cConsensus-based decentralized auctions for robust '
    'task allocation,\u201d IEEE Trans. Robotics, vol. 25, no. 4, pp. 912\u2013926, 2009.',
    'M. Ruediger and A. G. Banerjee, \u201cMapping-guided task discovery and allocation for cooperative '
    'robotic inspection of underwater structures (Map-TIDAL),\u201d arXiv:2602.02389, 2026.',
    '\u201cA centralized multi-robot task allocation for industrial plant inspection by using A* and '
    'genetic algorithms.\u201d',
    'E. Galceran and M. Carreras, \u201cA survey on coverage path planning for robotics,\u201d Robotics and '
    'Autonomous Systems, vol. 61, no. 12, pp. 1258\u20131276, 2013.',
    '\u201cA hierarchical LLM-based framework for heterogeneous multi-robot orchestration in high-risk '
    'energy facility maintenance,\u201d 2026.',
    'P. E. Hart, N. J. Nilsson, and B. Raphael, \u201cA formal basis for the heuristic determination of '
    'minimum cost paths,\u201d IEEE Trans. Systems Science and Cybernetics, vol. 4, no. 2, 1968.',
    'S. Macenski, T. Foote, B. Gerkey, C. Lalancette, and W. Woodall, \u201cRobot Operating System 2: '
    'design, architecture, and uses in the wild,\u201d Science Robotics, vol. 7, no. 66, 2022.',
    'N. Koenig and A. Howard, \u201cDesign and use paradigms for Gazebo, an open-source multi-robot '
    'simulator,\u201d in Proc. IEEE/RSJ IROS, 2004, pp. 2149\u20132154.',
]


def slide_references(prs):
    s = blank(prs)
    title(s, 'References')
    half = 7
    for col, (x, items) in enumerate([(0.60, REFERENCES[:half]), (6.75, REFERENCES[half:])]):
        y = 1.55
        for i, ref in enumerate(items, start=col * half + 1):
            bar(s, x, y + 0.09, 0.09, 0.09, DEEP)
            _, tf = textbox(s, x + 0.26, y, 5.70, 0.74)
            write(tf, [{'t': [{'t': f'[{i}]  ', 'bold': True, 'color': DEEP},
                              {'t': ref, 'color': INK}], 'size': 9.5}])
            y += 0.78
    _, tf = textbox(s, 0.60, 7.00, 12.10, 0.30)
    write(tf, [{'t': 'Entries [1]\u2013[5], [8], [9] and [11] are the project\u2019s domain sources; '
                     '[6], [7], [10], [12]\u2013[14] are the methodological references cited in the paper.',
                'size': 9, 'italic': True, 'color': MUTED}])
    return s


# ===========================================================================
#  Edits to the existing slides
# ===========================================================================
def set_title_text(slide, new_text):
    """Replaces the text of the title shape, preserving its run formatting."""
    for sh in slide.shapes:
        if sh.has_text_frame and sh.text_frame.paragraphs:
            runs = [r for para in sh.text_frame.paragraphs for r in para.runs]
            if runs and runs[0].font.size and runs[0].font.size.pt >= 30:
                runs[0].text = new_text
                for r in runs[1:]:
                    r.text = ''
                return True
    return False


def replace_run_text(slide, old_fragment, new_text):
    """Replaces the first run whose text starts with old_fragment."""
    for sh in slide.shapes:
        if not sh.has_text_frame:
            continue
        for para in sh.text_frame.paragraphs:
            for r in para.runs:
                if r.text.strip().startswith(old_fragment):
                    r.text = new_text
                    return True
    return False


RETITLE = {
    1: 'Introduction: Hull Fouling & Structural Defects',
    3: 'Proposed Work: Simulation Environment',
    4: 'Proposed Work: Mapping & Localization Layer',
    5: 'Literature Survey: Where the Open Gap Lies',
    6: 'Proposed Work: The Self-Verifying Cognitive Loop',
    7: 'Proposed Work: System Architecture',
    8: 'Literature Survey: Positioning vs. Published Work',
    10: 'Achieved Outcomes & Impact',
}

# Slide 11 was written before the system was evaluated; bring it to measured fact.
OUTCOME_EDITS = [
    ('A working belief-driven, severity-aware task prioritization layer',
     'A belief-driven, severity-aware prioritization layer with closed-loop verification, '
     'measured at 82.17% ordering correctness against 50.38% for an unprioritized baseline.'),
    ('Entire two-robot pipeline runs in Gazebo',
     'The full 14-node ROS 2 stack runs headless or with RViz2; 60 benchmark missions reproduce '
     'from a single command with fixed seeds.'),
    ('Directly benchmarked in concept against Map-TIDAL-style',
     'Benchmarked against a naive FIFO baseline sharing identical physics, defects and controllers, '
     'isolating the cognition layer: Spearman rₛ +0.70 vs −0.05.'),
    ('Belief map + utility scoring is detector-agnostic',
     'Belief map and utility scoring remain detector-agnostic — a learned segmentation model can '
     'replace the oracle detector without redesigning the cognition layer.'),
]


def main():
    prs = Presentation(DECK)
    n_orig = len(prs.slides)
    print(f'original slides: {n_orig}')

    for idx, text in RETITLE.items():
        ok = set_title_text(prs.slides[idx], text)
        print(f'  retitle [{idx}] -> {text[:48]!r} {"ok" if ok else "FAILED"}')

    for frag, new in OUTCOME_EDITS:
        ok = replace_run_text(prs.slides[10], frag, new)
        print(f'  outcome edit {frag[:34]!r} {"ok" if ok else "FAILED"}')

    builders = [slide_problem_statement, slide_utility, slide_results_walkthrough,
                slide_results_table, slide_results_charts, slide_discussion,
                slide_conclusion, slide_references]
    new_idx = {}
    for b in builders:
        b(prs)
        new_idx[b.__name__] = len(prs.slides) - 1
        print(f'  built {b.__name__} -> index {new_idx[b.__name__]}')

    order = [
        0,                                   # title
        1,                                   # introduction
        2,                                   # why simulation
        5,                                   # literature survey: gap
        8,                                   # literature survey: positioning
        new_idx['slide_problem_statement'],
        6,                                   # proposed: novelty loop
        7,                                   # proposed: architecture
        new_idx['slide_utility'],
        3,                                   # proposed: simulation environment
        4,                                   # proposed: mapping layer
        new_idx['slide_results_walkthrough'],
        new_idx['slide_results_table'],
        new_idx['slide_results_charts'],
        new_idx['slide_discussion'],
        10,                                  # achieved outcomes
        new_idx['slide_conclusion'],
        new_idx['slide_references'],
        11,                                  # thank you
    ]
    # The original references slide (index 9) is superseded and dropped.
    assert sorted(order) == [i for i in range(len(prs.slides)) if i != 9], \
        f'unexpected slide order: {sorted(order)}'

    sld_lst = prs.slides._sldIdLst
    ids = list(sld_lst)
    for el in ids:
        sld_lst.remove(el)
    for i in order:
        sld_lst.append(ids[i])
    # Drop the superseded references slide's part from the package.
    dropped = ids[9]
    rid = dropped.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id')
    prs.part.drop_rel(rid)
    print(f'reordered to {len(order)} slides')

    # Renumber: title and closing slides carry no number.
    for n, slide in enumerate(prs.slides, start=1):
        existing = [sh for sh in slide.shapes
                    if sh.has_text_frame and sh.text_frame.text.strip().isdigit()
                    and Emu(sh.left).inches > 12.0]
        if n in (1, len(order)):
            for sh in existing:
                sh._element.getparent().remove(sh._element)
            continue
        if existing:
            for para in existing[0].text_frame.paragraphs:
                for r in para.runs:
                    r.text = str(n)
                    break
                break
        else:
            pagenum(slide, n)

    prs.save(DECK)
    print(f'saved {DECK} with {len(prs.slides)} slides')


if __name__ == '__main__':
    main()
